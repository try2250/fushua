import openpyxl
from io import BytesIO
from tests.conftest import create_test_user, register_and_login, create_test_question
from app.models import ClassGroup, ClassMember, Assignment, AssignmentRecord, Record


class TestQuestionExportExcel:
    def test_export_questions_excel(self, client, db_session):
        teacher = create_test_user(db_session, "excelteacher", "teacher")
        create_test_question(db_session, subject="数学", created_by=teacher.id, content="1+1=?")
        create_test_question(db_session, subject="语文", created_by=teacher.id, content="默写古诗")
        register_and_login(client, "excelteacher", "teacher")
        response = client.get("/teacher/questions/export/excel")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers.get("content-type", "")
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb.active
        assert ws.title == "题库"
        assert ws.cell(row=1, column=1).value == "ID"
        assert ws.max_row == 3

    def test_export_questions_excel_empty(self, client, db_session):
        register_and_login(client, "excelteacher2", "teacher")
        response = client.get("/teacher/questions/export/excel")
        assert response.status_code == 200
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb.active
        assert ws.max_row == 1

    def test_export_questions_excel_requires_login(self, client, db_session):
        response = client.get("/teacher/questions/export/excel")
        assert response.status_code in (303, 403)

    def test_export_questions_excel_requires_teacher(self, client, db_session):
        register_and_login(client, "excelstudent", "student")
        response = client.get("/teacher/questions/export/excel")
        assert response.status_code in (303, 403)


class TestClassExportExcel:
    def test_export_class_excel(self, client, db_session):
        teacher = create_test_user(db_session, "classexcelteacher", "teacher")
        student = create_test_user(db_session, "classexcelstudent", "student")
        cls = ClassGroup(name="测试班", created_by=teacher.id)
        db_session.add(cls)
        db_session.commit()
        db_session.refresh(cls)
        db_session.add(ClassMember(class_id=cls.id, user_id=student.id))
        db_session.commit()
        q = create_test_question(db_session, created_by=teacher.id)
        db_session.add(Record(user_id=student.id, question_id=q.id, user_answer="B", is_correct=True))
        db_session.commit()
        register_and_login(client, "classexcelteacher", "teacher")
        response = client.get(f"/teacher/classes/{cls.id}/export/excel")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers.get("content-type", "")
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb.active
        assert ws.title == "班级报告"
        assert ws.cell(row=1, column=1).value == "排名"

    def test_export_class_excel_not_found(self, client, db_session):
        register_and_login(client, "classexcelteacher2", "teacher")
        response = client.get("/teacher/classes/99999/export/excel")
        assert response.status_code == 404

    def test_export_class_excel_requires_teacher(self, client, db_session):
        register_and_login(client, "classexcelstudent2", "student")
        response = client.get("/teacher/classes/1/export/excel")
        assert response.status_code in (303, 403)


class TestAssignmentExportExcel:
    def test_export_assignment_excel(self, client, db_session):
        teacher = create_test_user(db_session, "asgnexcelteacher", "teacher")
        student = create_test_user(db_session, "asgnexcelstudent", "student")
        q = create_test_question(db_session, created_by=teacher.id)
        cls = ClassGroup(name="作业班", created_by=teacher.id)
        db_session.add(cls)
        db_session.commit()
        db_session.refresh(cls)
        db_session.add(ClassMember(class_id=cls.id, user_id=student.id))
        db_session.commit()
        assignment = Assignment(
            title="测试作业",
            question_ids=str(q.id),
            created_by=teacher.id,
            class_id=cls.id,
        )
        db_session.add(assignment)
        db_session.commit()
        db_session.refresh(assignment)
        db_session.add(AssignmentRecord(assignment_id=assignment.id, user_id=student.id, completed=True))
        db_session.add(Record(user_id=student.id, question_id=q.id, user_answer="B", is_correct=True))
        db_session.commit()
        register_and_login(client, "asgnexcelteacher", "teacher")
        response = client.get(f"/teacher/assignments/{assignment.id}/export/excel")
        assert response.status_code == 200
        assert "spreadsheetml" in response.headers.get("content-type", "")
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb.active
        assert ws.title == "作业完成情况"
        assert ws.cell(row=1, column=1).value == "姓名"
        assert ws.max_row >= 2

    def test_export_assignment_excel_not_found(self, client, db_session):
        register_and_login(client, "asgnexcelteacher2", "teacher")
        response = client.get("/teacher/assignments/99999/export/excel")
        assert response.status_code == 404

    def test_export_assignment_excel_requires_teacher(self, client, db_session):
        register_and_login(client, "asgnexcelstudent3", "student")
        response = client.get("/teacher/assignments/1/export/excel")
        assert response.status_code in (303, 403)

    def test_export_assignment_excel_incomplete_student(self, client, db_session):
        teacher = create_test_user(db_session, "asgnexcelteacher3", "teacher")
        student = create_test_user(db_session, "asgnexcelstudent4", "student")
        q = create_test_question(db_session, created_by=teacher.id)
        cls = ClassGroup(name="作业班2", created_by=teacher.id)
        db_session.add(cls)
        db_session.commit()
        db_session.refresh(cls)
        db_session.add(ClassMember(class_id=cls.id, user_id=student.id))
        db_session.commit()
        assignment = Assignment(
            title="未完成作业",
            question_ids=str(q.id),
            created_by=teacher.id,
            class_id=cls.id,
        )
        db_session.add(assignment)
        db_session.commit()
        db_session.refresh(assignment)
        register_and_login(client, "asgnexcelteacher3", "teacher")
        response = client.get(f"/teacher/assignments/{assignment.id}/export/excel")
        assert response.status_code == 200
        wb = openpyxl.load_workbook(BytesIO(response.content))
        ws = wb.active
        found_incomplete = False
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[2] == "未完成":
                found_incomplete = True
                break
        assert found_incomplete
