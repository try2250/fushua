import json
import csv
import io
import uuid
import urllib.parse
from fastapi import APIRouter, Depends, Request, Form, UploadFile, File, HTTPException
from fastapi.responses import RedirectResponse, Response, StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func as sa_func, Integer
from typing import Annotated
from io import BytesIO
from datetime import datetime, timedelta

import openpyxl

from app.database import get_db
from app.models import Question, User, Record, FieldConfig, QuestionBank, SiteConfig, ClassGroup, ClassMember, Notification, Favorite, Assignment, AssignmentRecord, ClassJoinRequest, QUESTION_TYPES, SEMESTERS, SUBJECTS, BUILTIN_FIELDS, FIELD_TYPE_CHOICES, AuditLog
from app.auth import require_teacher, require_admin_role, get_current_user
from app.routers.permissions import is_admin, teacher_owns_bank, teacher_owns_student
from app.security import validate_csrf_async, sanitize_input
from app.utils.validation import parse_int, paginate

router = APIRouter(prefix="/teacher")

MAX_UPLOAD_SIZE = 5 * 1024 * 1024

_pending_imports = {}


def _parse_owned_bank_id(db: Session, teacher_id: int, raw_bank_id) -> int | None:
    if not raw_bank_id:
        return None
    bank_id = parse_int(raw_bank_id, min_value=1)
    if bank_id is None:
        raise HTTPException(status_code=400, detail="Invalid bank id")
    if not teacher_owns_bank(db, teacher_id, bank_id):
        raise HTTPException(status_code=404, detail="Question bank not found")
    return bank_id


def require_admin(request: Request, db: Session) -> int:
    from app.auth import require_admin_role
    return require_admin_role(request, db)


@router.get("/questions")
def manage_questions(
    request: Request,
    subject: str = "",
    q_type: str = "",
    page: str = "1",
    db: Annotated[Session, Depends(get_db)] = None,
):
    user_id = require_teacher(request, db)
    query = db.query(Question).filter(Question.created_by == user_id)
    if subject:
        query = query.filter(Question.subject == subject)
    if q_type:
        query = query.filter(Question.q_type == q_type)
    query = query.order_by(Question.created_at.desc())
    page_num = parse_int(page, default=1, min_value=1) or 1
    pagination = paginate(query, page_num, per_page=20)
    subjects = [s[0] for s in db.query(Question.subject).distinct().all()]
    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
    query_params = []
    if subject:
        query_params.append(f"subject={subject}")
    if q_type:
        query_params.append(f"q_type={q_type}")
    return request.app.state.templates.TemplateResponse(
        "teacher/questions.html",
        {
            "request": request,
            "questions": pagination["items"],
            "subjects": subjects,
            "current_subject": subject,
            "current_type": q_type,
            "question_types": QUESTION_TYPES,
            "semesters": SEMESTERS,
            "custom_fields": custom_fields,
            "csrf_token": request.session.get("csrf_token", ""),
            "pagination": pagination,
            "base_url": "/teacher/questions?",
            "query_params": "&".join(query_params),
        },
    )


@router.get("/questions/create")
def create_question_page(request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
    banks = db.query(QuestionBank).filter(QuestionBank.created_by == user_id).order_by(QuestionBank.name).all()
    return request.app.state.templates.TemplateResponse(
        "teacher/question_form.html",
        {
            "request": request,
            "question": None,
            "error": None,
            "question_types": QUESTION_TYPES,
            "semesters": SEMESTERS,
            "custom_fields": custom_fields,
            "banks": banks,
            "csrf_token": request.session.get("csrf_token", ""),
        },
    )


@router.post("/questions/create")
async def create_question(
    request: Request,
    db: Annotated[Session, Depends(get_db)] = None,
):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    form = await request.form()
    bank_id_val = form.get("bank_id", "")
    subject = form.get("subject", "")
    content = form.get("content", "")
    answer = form.get("answer", "")
    subject = sanitize_input(subject, max_length=20)
    content = sanitize_input(content, max_length=5000)
    answer = sanitize_input(answer, max_length=200)
    if not subject or not content or not answer:
        custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
        return request.app.state.templates.TemplateResponse(
            "teacher/question_form.html",
            {
                "request": request,
                "question": None,
                "error": "科目、题目内容和答案为必填项",
                "question_types": QUESTION_TYPES,
                "semesters": SEMESTERS,
                "custom_fields": custom_fields,
                "csrf_token": request.session.get("csrf_token", ""),
            },
        )

    bank_id = _parse_owned_bank_id(db, user_id, bank_id_val)

    question = Question(
        subject=subject,
        semester=form.get("semester", ""),
        chapter=form.get("chapter", ""),
        difficulty=parse_int(form.get("difficulty"), default=2, min_value=1, max_value=5) or 2,
        q_type=form.get("q_type", "choice"),
        content=content,
        option_a=form.get("option_a", ""),
        option_b=form.get("option_b", ""),
        option_c=form.get("option_c", ""),
        option_d=form.get("option_d", ""),
        answer=answer,
        explanation=form.get("explanation", ""),
        image_url=form.get("image_url", ""),
        bank_id=bank_id,
        created_by=user_id,
    )

    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).all()
    extra = {}
    for cf in custom_fields:
        val = form.get(f"extra_{cf.field_key}", "")
        if val:
            extra[cf.field_key] = val
    question.extra_data = json.dumps(extra, ensure_ascii=False)

    db.add(question)
    db.commit()

    if form.get("continue_creating"):
        return RedirectResponse(url="/teacher/questions/create", status_code=303)

    return RedirectResponse(url="/teacher/questions", status_code=303)


@router.get("/questions/{question_id}/edit")
def edit_question_page(question_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    question = db.query(Question).filter(
        Question.id == question_id, Question.created_by == user_id
    ).first()
    if not question:
        raise HTTPException(status_code=404, detail="题目不存在")
    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
    banks = db.query(QuestionBank).filter(QuestionBank.created_by == user_id).order_by(QuestionBank.name).all()
    return request.app.state.templates.TemplateResponse(
        "teacher/question_form.html",
        {
            "request": request,
            "question": question,
            "error": None,
            "question_types": QUESTION_TYPES,
            "semesters": SEMESTERS,
            "custom_fields": custom_fields,
            "banks": banks,
            "csrf_token": request.session.get("csrf_token", ""),
        },
    )


@router.post("/questions/{question_id}/edit")
async def edit_question(
    question_id: int,
    request: Request,
    db: Annotated[Session, Depends(get_db)] = None,
):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    question = db.query(Question).filter(
        Question.id == question_id, Question.created_by == user_id
    ).first()
    if not question:
        raise HTTPException(status_code=404, detail="题目不存在")

    form = await request.form()
    bank_id_val = form.get("bank_id", "")
    question.bank_id = _parse_owned_bank_id(db, user_id, bank_id_val)
    question.subject = form.get("subject", question.subject)
    question.semester = form.get("semester", "")
    question.chapter = form.get("chapter", "")
    question.difficulty = parse_int(form.get("difficulty"), default=2, min_value=1, max_value=5) or 2
    question.q_type = form.get("q_type", "choice")
    question.content = form.get("content", question.content)
    question.option_a = form.get("option_a", "")
    question.option_b = form.get("option_b", "")
    question.option_c = form.get("option_c", "")
    question.option_d = form.get("option_d", "")
    question.answer = form.get("answer", question.answer)
    question.explanation = form.get("explanation", "")
    question.image_url = form.get("image_url", "")

    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).all()
    extra = question.extra
    for cf in custom_fields:
        val = form.get(f"extra_{cf.field_key}", "")
        extra[cf.field_key] = val
    question.extra_data = json.dumps(extra, ensure_ascii=False)

    db.commit()
    return RedirectResponse(url="/teacher/questions", status_code=303)


@router.get("/banks")
def bank_list(request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    banks = db.query(QuestionBank).filter(QuestionBank.created_by == user_id).order_by(QuestionBank.created_at.desc()).all()
    bank_data = []
    for b in banks:
        q_count = db.query(Question).filter(Question.bank_id == b.id).count()
        bank_data.append({"id": b.id, "name": b.name, "subject": b.subject, "semester": b.semester, "bank_type": b.bank_type, "visibility": b.visibility, "access_code": b.access_code, "description": b.description, "q_count": q_count, "created_at": b.created_at})
    return request.app.state.templates.TemplateResponse(
        "teacher/banks.html",
        {"request": request, "banks": bank_data, "subjects": SUBJECTS, "semesters": SEMESTERS, "csrf_token": request.session.get("csrf_token", "")},
    )


@router.get("/banks/create")
def create_bank_page(request: Request, db: Annotated[Session, Depends(get_db)]):
    require_teacher(request, db)
    return request.app.state.templates.TemplateResponse(
        "teacher/bank_form.html",
        {"request": request, "bank": None, "subjects": SUBJECTS, "semesters": SEMESTERS, "csrf_token": request.session.get("csrf_token", "")},
    )


@router.post("/banks/create")
async def create_bank(request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    form = await request.form()
    name = sanitize_input(form.get("name", "").strip(), max_length=100)
    subject = sanitize_input(form.get("subject", "").strip(), max_length=20)
    if not name or not subject:
        return request.app.state.templates.TemplateResponse(
            "teacher/bank_form.html",
            {"request": request, "bank": None, "error": "题库名称和科目为必填项", "subjects": SUBJECTS, "semesters": SEMESTERS, "csrf_token": request.session.get("csrf_token", "")},
        )
    bank = QuestionBank(
        name=name,
        subject=subject,
        semester=form.get("semester", ""),
        description=sanitize_input(form.get("description", "").strip(), max_length=500),
        bank_type=form.get("bank_type", "custom"),
        visibility=form.get("visibility", "public"),
        access_code=sanitize_input(form.get("access_code", "").strip(), max_length=50) if form.get("visibility") == "code" else "",
        created_by=user_id,
    )
    db.add(bank)
    db.commit()
    return RedirectResponse(url="/teacher/banks", status_code=303)


@router.post("/banks/{bank_id}/delete")
async def delete_bank(bank_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    bank = db.query(QuestionBank).filter(QuestionBank.id == bank_id).first()
    if not bank or not teacher_owns_bank(db, user_id, bank_id):
        raise HTTPException(status_code=404)
    db.add(AuditLog(
        actor_id=user_id,
        action="delete_bank",
        target_type="question_bank",
        target_id=bank_id,
        detail=f"删除题库「{bank.name}」"
    ))
    db.query(Question).filter(Question.bank_id == bank_id).update({"bank_id": None})
    db.delete(bank)
    db.commit()
    return RedirectResponse(url="/teacher/banks", status_code=303)


@router.get("/fields")
def field_manager(request: Request, db: Annotated[Session, Depends(get_db)]):
    require_teacher(request, db)
    fields = db.query(FieldConfig).order_by(FieldConfig.sort_order, FieldConfig.id).all()
    return request.app.state.templates.TemplateResponse(
        "teacher/fields.html",
        {
            "request": request,
            "fields": fields,
            "builtin_fields": BUILTIN_FIELDS,
            "field_type_choices": FIELD_TYPE_CHOICES,
            "csrf_token": request.session.get("csrf_token", ""),
        },
    )


@router.post("/fields/create")
async def create_field(request: Request, db: Annotated[Session, Depends(get_db)]):
    require_teacher(request, db)
    await validate_csrf_async(request)
    form = await request.form()
    field_key = form.get("field_key", "").strip()
    field_label = form.get("field_label", "").strip()
    field_key = sanitize_input(field_key, max_length=50)
    field_label = sanitize_input(field_label, max_length=100)
    field_type = form.get("field_type", "text")
    required = form.get("required") == "on"
    visible = form.get("visible") == "on"
    options = form.get("options", "").strip()
    sort_order = parse_int(form.get("sort_order"), default=0) or 0

    if not field_key or not field_label:
        return RedirectResponse(url="/teacher/fields", status_code=303)

    if field_key in BUILTIN_FIELDS:
        return RedirectResponse(url="/teacher/fields", status_code=303)

    existing = db.query(FieldConfig).filter(FieldConfig.field_key == field_key).first()
    if existing:
        return RedirectResponse(url="/teacher/fields", status_code=303)

    fc = FieldConfig(
        field_key=field_key,
        field_label=field_label,
        field_type=field_type,
        required=required,
        visible=visible,
        options=options,
        sort_order=sort_order,
    )
    db.add(fc)
    db.commit()
    return RedirectResponse(url="/teacher/fields", status_code=303)


@router.post("/fields/{field_id}/update")
async def update_field(field_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    require_teacher(request, db)
    await validate_csrf_async(request)
    fc = db.query(FieldConfig).filter(FieldConfig.id == field_id).first()
    if not fc:
        raise HTTPException(status_code=404, detail="字段不存在")

    form = await request.form()
    fc.field_label = form.get("field_label", fc.field_label)
    fc.field_type = form.get("field_type", fc.field_type)
    fc.required = form.get("required") == "on"
    fc.visible = form.get("visible") == "on"
    fc.options = form.get("options", "")
    fc.sort_order = parse_int(form.get("sort_order"), default=fc.sort_order) or fc.sort_order
    db.commit()
    return RedirectResponse(url="/teacher/fields", status_code=303)


@router.post("/fields/{field_id}/delete")
async def delete_field(field_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    require_teacher(request, db)
    await validate_csrf_async(request)
    fc = db.query(FieldConfig).filter(FieldConfig.id == field_id).first()
    if not fc:
        raise HTTPException(status_code=404, detail="字段不存在")
    db.delete(fc)
    db.commit()
    return RedirectResponse(url="/teacher/fields", status_code=303)


@router.get("/questions/import/template/{fmt}")
def download_template(fmt: str, request: Request, db: Annotated[Session, Depends(get_db)]):
    require_teacher(request, db)
    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
    custom_keys = [cf.field_key for cf in custom_fields]

    if fmt == "json":
        example = [{
            "subject": "数学",
            "semester": "七年级上册",
            "chapter": "有理数",
            "q_type": "choice",
            "difficulty": 2,
            "content": "题目内容",
            "option_a": "选项A",
            "option_b": "选项B",
            "option_c": "选项C",
            "option_d": "选项D",
            "answer": "A",
            "explanation": "解析说明",
        }]
        for ck in custom_keys:
            example[0][ck] = f"自定义字段{ck}的值"
        return Response(
            content=json.dumps(example, ensure_ascii=False, indent=2),
            media_type="application/json",
            headers={"Content-Disposition": "attachment; filename=import_template.json"},
        )
    elif fmt == "csv":
        fieldnames = ["subject", "semester", "chapter", "q_type", "difficulty", "content", "option_a", "option_b", "option_c", "option_d", "answer", "explanation"] + custom_keys
        example_row = {"subject": "数学", "semester": "七年级上册", "chapter": "有理数", "q_type": "choice", "difficulty": "2", "content": "题目内容", "option_a": "选项A", "option_b": "选项B", "option_c": "选项C", "option_d": "选项D", "answer": "A", "explanation": "解析说明"}
        for ck in custom_keys:
            example_row[ck] = f"自定义字段{ck}的值"
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerow(example_row)
        return Response(
            content="\ufeff" + output.getvalue(),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": "attachment; filename=import_template.csv"},
        )
    raise HTTPException(status_code=400, detail="不支持的格式")


@router.get("/questions/import")
def import_page(request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
    banks = db.query(QuestionBank).filter(QuestionBank.created_by == user_id).order_by(QuestionBank.name).all()
    return request.app.state.templates.TemplateResponse(
        "teacher/import.html",
        {
            "request": request,
            "error": None,
            "success": None,
            "question_types": QUESTION_TYPES,
            "custom_fields": custom_fields,
            "banks": banks,
            "csrf_token": request.session.get("csrf_token", ""),
        },
    )


@router.post("/questions/import")
async def import_questions(
    request: Request,
    file: UploadFile = File(...),
    db: Annotated[Session, Depends(get_db)] = None,
):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    form_data = await request.form()
    bank_id = _parse_owned_bank_id(db, user_id, form_data.get("bank_id", ""))
    filename = file.filename or ""
    content_bytes = await file.read()

    if len(content_bytes) > MAX_UPLOAD_SIZE:
        custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
        return request.app.state.templates.TemplateResponse(
            "teacher/import.html",
            {
                "request": request,
                "error": f"文件大小超过限制（最大 {MAX_UPLOAD_SIZE // 1024 // 1024}MB）",
                "success": None,
                "question_types": QUESTION_TYPES,
                "custom_fields": custom_fields,
                "csrf_token": request.session.get("csrf_token", ""),
            },
        )

    try:
        if filename.endswith(".json"):
            count = _import_json(content_bytes, user_id, db, bank_id=bank_id)
        elif filename.endswith(".csv"):
            count = _import_csv(content_bytes, user_id, db, bank_id=bank_id)
        else:
            custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
            return request.app.state.templates.TemplateResponse(
                "teacher/import.html",
                {
                    "request": request,
                    "error": "仅支持 JSON 和 CSV 格式",
                    "success": None,
                    "question_types": QUESTION_TYPES,
                    "custom_fields": custom_fields,
                    "csrf_token": request.session.get("csrf_token", ""),
                },
            )
    except Exception as e:
        custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
        return request.app.state.templates.TemplateResponse(
            "teacher/import.html",
            {
                "request": request,
                "error": "导入失败，请检查文件格式是否正确",
                "success": None,
                "question_types": QUESTION_TYPES,
                "custom_fields": custom_fields,
                "csrf_token": request.session.get("csrf_token", ""),
            },
        )

    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
    return request.app.state.templates.TemplateResponse(
        "teacher/import.html",
        {
            "request": request,
            "error": None,
            "success": f"成功导入 {count} 道题目！",
            "question_types": QUESTION_TYPES,
            "custom_fields": custom_fields,
            "csrf_token": request.session.get("csrf_token", ""),
        },
    )


def _parse_file_content(content_bytes: bytes, filename: str):
    rows = []
    if filename.endswith(".json"):
        data = json.loads(content_bytes.decode("utf-8"))
        if not isinstance(data, list):
            data = [data]
        for idx, item in enumerate(data):
            row = {k: (v.strip() if isinstance(v, str) else str(v)) for k, v in item.items()}
            row["_row_num"] = idx + 1
            rows.append(row)
    elif filename.endswith(".csv"):
        text = content_bytes.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        for idx, raw in enumerate(reader):
            row = {k: (v.strip() if isinstance(v, str) else "") for k, v in raw.items() if k is not None}
            row["_row_num"] = idx + 1
            rows.append(row)
    return rows


def _validate_rows(rows):
    error_rows = []
    valid_rows = []
    for row in rows:
        missing = []
        if not row.get("subject"):
            missing.append("subject")
        if not row.get("content"):
            missing.append("content")
        if not row.get("answer"):
            missing.append("answer")
        if missing:
            error_rows.append({"row_num": row["_row_num"], "missing_fields": missing, "data": row})
        else:
            valid_rows.append(row)
    return valid_rows, error_rows


def _check_duplicates(valid_rows, user_id, db):
    duplicate_rows = []
    unique_rows = []
    existing_pairs = set()
    questions = db.query(Question.content, Question.answer).filter(Question.created_by == user_id).all()
    for q in questions:
        existing_pairs.add((q.content, q.answer))

    seen_in_file = set()
    for row in valid_rows:
        key = (row.get("content", ""), row.get("answer", ""))
        if key in existing_pairs or key in seen_in_file:
            duplicate_rows.append({"row_num": row["_row_num"], "content": row.get("content", ""), "answer": row.get("answer", "")})
        else:
            seen_in_file.add(key)
            unique_rows.append(row)
    return unique_rows, duplicate_rows


@router.post("/questions/import-preview")
async def import_preview(
    request: Request,
    file: UploadFile = File(...),
    db: Annotated[Session, Depends(get_db)] = None,
):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    form_data = await request.form()
    bank_id = _parse_owned_bank_id(db, user_id, form_data.get("bank_id", ""))
    filename = file.filename or ""
    content_bytes = await file.read()

    if len(content_bytes) > MAX_UPLOAD_SIZE:
        custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
        banks = db.query(QuestionBank).filter(QuestionBank.created_by == user_id).order_by(QuestionBank.name).all()
        return request.app.state.templates.TemplateResponse(
            "teacher/import.html",
            {
                "request": request,
                "error": f"文件大小超过限制（最大 {MAX_UPLOAD_SIZE // 1024 // 1024}MB）",
                "success": None,
                "question_types": QUESTION_TYPES,
                "custom_fields": custom_fields,
                "banks": banks,
                "csrf_token": request.session.get("csrf_token", ""),
            },
        )

    if not filename.endswith(".json") and not filename.endswith(".csv"):
        custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
        banks = db.query(QuestionBank).filter(QuestionBank.created_by == user_id).order_by(QuestionBank.name).all()
        return request.app.state.templates.TemplateResponse(
            "teacher/import.html",
            {
                "request": request,
                "error": "仅支持 JSON 和 CSV 格式",
                "success": None,
                "question_types": QUESTION_TYPES,
                "custom_fields": custom_fields,
                "banks": banks,
                "csrf_token": request.session.get("csrf_token", ""),
            },
        )

    try:
        rows = _parse_file_content(content_bytes, filename)
    except Exception as e:
        custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
        banks = db.query(QuestionBank).filter(QuestionBank.created_by == user_id).order_by(QuestionBank.name).all()
        return request.app.state.templates.TemplateResponse(
            "teacher/import.html",
            {
                "request": request,
                "error": "文件解析失败，请检查文件格式是否正确",
                "success": None,
                "question_types": QUESTION_TYPES,
                "custom_fields": custom_fields,
                "banks": banks,
                "csrf_token": request.session.get("csrf_token", ""),
            },
        )

    valid_rows, error_rows = _validate_rows(rows)
    unique_rows, duplicate_rows = _check_duplicates(valid_rows, user_id, db)

    preview_rows = unique_rows[:20]
    total_rows = len(rows)
    valid_count = len(unique_rows)

    import_token = uuid.uuid4().hex
    _pending_imports[import_token] = {
        "rows": unique_rows,
        "user_id": user_id,
        "bank_id": bank_id,
        "filename": filename,
    }

    return request.app.state.templates.TemplateResponse(
        "teacher/import_preview.html",
        {
            "request": request,
            "preview_rows": preview_rows,
            "error_rows": error_rows,
            "duplicate_rows": duplicate_rows,
            "total_rows": total_rows,
            "valid_count": valid_count,
            "error_count": len(error_rows),
            "duplicate_count": len(duplicate_rows),
            "import_token": import_token,
            "bank_id": bank_id,
            "filename": filename,
            "csrf_token": request.session.get("csrf_token", ""),
        },
    )


@router.post("/questions/import-confirm")
async def import_confirm(request: Request, db: Annotated[Session, Depends(get_db)] = None):
    user_id = require_teacher(request, db)
    form = await request.form()
    import_token = form.get("import_token", "")
    await validate_csrf_async(request)

    pending = _pending_imports.pop(import_token, None)
    if not pending or pending["user_id"] != user_id:
        custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
        banks = db.query(QuestionBank).filter(QuestionBank.created_by == user_id).order_by(QuestionBank.name).all()
        return request.app.state.templates.TemplateResponse(
            "teacher/import.html",
            {
                "request": request,
                "error": "导入会话已过期，请重新上传文件",
                "success": None,
                "question_types": QUESTION_TYPES,
                "custom_fields": custom_fields,
                "banks": banks,
                "csrf_token": request.session.get("csrf_token", ""),
            },
        )

    rows = pending["rows"]
    bank_id = pending["bank_id"]
    count = 0
    batch_size = 50
    for row in rows:
        item = {k: v for k, v in row.items() if k != "_row_num"}
        q = _build_question_from_dict(item, user_id, bank_id=bank_id)
        db.add(q)
        count += 1
        if count > 0 and count % batch_size == 0:
            db.commit()
    db.add(AuditLog(
        actor_id=user_id,
        action="import_questions",
        target_type="question",
        detail=f"批量导入 {count} 道题目，题库ID={bank_id or '无'}"
    ))
    db.commit()
    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
    banks = db.query(QuestionBank).filter(QuestionBank.created_by == user_id).order_by(QuestionBank.name).all()
    return request.app.state.templates.TemplateResponse(
        "teacher/import.html",
        {
            "request": request,
            "error": None,
            "success": f"成功导入 {count} 道题目！",
            "question_types": QUESTION_TYPES,
            "custom_fields": custom_fields,
            "banks": banks,
            "csrf_token": request.session.get("csrf_token", ""),
        },
    )


def _build_question_from_dict(item: dict, user_id: int, bank_id: int = None) -> Question:
    builtin_data = {}
    extra_data = {}
    all_builtin = set(BUILTIN_FIELDS)
    for k, v in item.items():
        if k in all_builtin:
            builtin_data[k] = v
        else:
            extra_data[k] = v

    q = Question(
        subject=builtin_data.get("subject", ""),
        semester=builtin_data.get("semester", ""),
        chapter=builtin_data.get("chapter", ""),
        difficulty=parse_int(builtin_data.get("difficulty"), default=2, min_value=1, max_value=5) or 2,
        q_type=builtin_data.get("q_type", "choice"),
        content=builtin_data.get("content", ""),
        option_a=builtin_data.get("option_a", ""),
        option_b=builtin_data.get("option_b", ""),
        option_c=builtin_data.get("option_c", ""),
        option_d=builtin_data.get("option_d", ""),
        answer=builtin_data.get("answer", ""),
        explanation=builtin_data.get("explanation", ""),
        extra_data=json.dumps(extra_data, ensure_ascii=False) if extra_data else "{}",
        bank_id=bank_id,
        created_by=user_id,
    )
    return q


def _import_json(content_bytes: bytes, user_id: int, db: Session, bank_id: int = None) -> int:
    data = json.loads(content_bytes.decode("utf-8"))
    if not isinstance(data, list):
        data = [data]
    count = 0
    batch_size = 50
    for i, item in enumerate(data):
        q = _build_question_from_dict(item, user_id, bank_id=bank_id)
        if q.subject and q.content and q.answer:
            db.add(q)
            count += 1
        if count > 0 and count % batch_size == 0:
            db.commit()
    db.commit()
    return count


def _import_csv(content_bytes: bytes, user_id: int, db: Session, bank_id: int = None) -> int:
    text = content_bytes.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    count = 0
    batch_size = 50
    for row in reader:
        item = {k: v.strip() for k, v in row.items() if v is not None}
        q = _build_question_from_dict(item, user_id, bank_id=bank_id)
        if q.subject and q.content and q.answer:
            db.add(q)
            count += 1
        if count > 0 and count % batch_size == 0:
            db.commit()
    db.commit()
    return count


@router.get("/questions/export/{fmt}")
def export_questions(fmt: str, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    questions = db.query(Question).filter(Question.created_by == user_id).all()
    custom_fields = db.query(FieldConfig).filter(FieldConfig.visible == True).order_by(FieldConfig.sort_order).all()
    custom_keys = [cf.field_key for cf in custom_fields]

    if fmt == "json":
        data = []
        for q in questions:
            item = {
                "subject": q.subject,
                "semester": q.semester,
                "chapter": q.chapter,
                "q_type": q.q_type,
                "difficulty": q.difficulty,
                "content": q.content,
                "option_a": q.option_a,
                "option_b": q.option_b,
                "option_c": q.option_c,
                "option_d": q.option_d,
                "answer": q.answer,
                "explanation": q.explanation,
            }
            extra = q.extra
            for ck in custom_keys:
                if ck in extra:
                    item[ck] = extra[ck]
            data.append(item)
        return Response(
            content=json.dumps(data, ensure_ascii=False, indent=2),
            media_type="application/json",
            headers={"Content-Disposition": "attachment; filename=questions.json"},
        )
    elif fmt == "csv":
        fieldnames = [
            "subject", "semester", "chapter", "q_type", "difficulty", "content",
            "option_a", "option_b", "option_c", "option_d", "answer", "explanation"
        ] + custom_keys
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        for q in questions:
            row = {
                "subject": q.subject,
                "semester": q.semester,
                "chapter": q.chapter,
                "q_type": q.q_type,
                "difficulty": q.difficulty,
                "content": q.content,
                "option_a": q.option_a,
                "option_b": q.option_b,
                "option_c": q.option_c,
                "option_d": q.option_d,
                "answer": q.answer,
                "explanation": q.explanation,
            }
            extra = q.extra
            for ck in custom_keys:
                row[ck] = extra.get(ck, "")
            writer.writerow(row)
        return Response(
            content="\ufeff" + output.getvalue(),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": "attachment; filename=questions.csv"},
        )
    elif fmt == "excel":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "题库"
        headers = ["ID", "科目", "学期", "章节", "题型", "难度", "题目内容", "选项A", "选项B", "选项C", "选项D", "答案", "解析"]
        for ck in custom_keys:
            headers.append(ck)
        ws.append(headers)
        for q in questions:
            row = [q.id, q.subject, q.semester, q.chapter, QUESTION_TYPES.get(q.q_type, q.q_type), q.difficulty, q.content, q.option_a, q.option_b, q.option_c, q.option_d, q.answer, q.explanation]
            extra = q.extra
            for ck in custom_keys:
                row.append(extra.get(ck, ""))
            ws.append(row)
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        filename = urllib.parse.quote("题库.xlsx")
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})
    raise HTTPException(status_code=400, detail="不支持的导出格式")


@router.post("/questions/batch-edit")
async def batch_edit_questions(request: Request, db: Annotated[Session, Depends(get_db)] = None):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    form = await request.form()
    question_ids_str = form.get("question_ids", "")
    action = form.get("action", "")
    value = form.get("value", "")

    if not question_ids_str or not action:
        raise HTTPException(status_code=400, detail="参数不完整")

    qids = []
    for raw_qid in [x.strip() for x in question_ids_str.split(",") if x.strip()]:
        qid = parse_int(raw_qid, min_value=1)
        if qid is None:
            raise HTTPException(status_code=400, detail="题目ID格式错误")
        qids.append(qid)

    if not qids:
        raise HTTPException(status_code=400, detail="未选择题目")

    questions = db.query(Question).filter(
        Question.id.in_(qids), Question.created_by == user_id
    ).all()

    if action == "delete":
        for q in questions:
            db.query(Record).filter(Record.question_id == q.id).delete()
            db.query(Favorite).filter(Favorite.question_id == q.id).delete()
            affected_assignments = db.query(Assignment).filter(
                Assignment.question_ids.contains(str(q.id))
            ).all()
            for a in affected_assignments:
                ids = [qid.strip() for qid in a.question_ids.split(",") if qid.strip() and qid.strip() != str(q.id)]
                a.question_ids = ",".join(ids)
            db.delete(q)
    elif action == "difficulty":
        diff = parse_int(value, min_value=1, max_value=5)
        if diff is None:
            raise HTTPException(status_code=400, detail="难度值必须为1-5的数字")
        for q in questions:
            q.difficulty = diff
    elif action == "semester":
        value = sanitize_input(value, max_length=20)
        for q in questions:
            q.semester = value
    elif action == "chapter":
        value = sanitize_input(value, max_length=100)
        for q in questions:
            q.chapter = value
    elif action == "bank":
        bank_id = _parse_owned_bank_id(db, user_id, value)
        for q in questions:
            q.bank_id = bank_id
    else:
        raise HTTPException(status_code=400, detail="不支持的操作类型")

    db.add(AuditLog(
        actor_id=user_id,
        action="batch_edit",
        target_type="question",
        detail=f"批量操作：{action}，影响 {len(questions)} 道题目"
    ))
    db.commit()

    return RedirectResponse(url="/teacher/questions", status_code=303)


@router.post("/questions/{question_id}/delete")
async def delete_question(question_id: int, request: Request, db: Annotated[Session, Depends(get_db)] = None):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    question = db.query(Question).filter(
        Question.id == question_id, Question.created_by == user_id
    ).first()
    if not question:
        raise HTTPException(status_code=404, detail="题目不存在")
    db.add(AuditLog(
        actor_id=user_id,
        action="delete_question",
        target_type="question",
        target_id=question_id,
        detail=f"删除题目「{question.content[:30]}」"
    ))
    db.query(Record).filter(Record.question_id == question_id).delete()
    db.query(Favorite).filter(Favorite.question_id == question_id).delete()
    affected_assignments = db.query(Assignment).filter(Assignment.question_ids.contains(str(question_id))).all()
    for a in affected_assignments:
        ids = [qid.strip() for qid in a.question_ids.split(",") if qid.strip() and qid.strip() != str(question_id)]
        a.question_ids = ",".join(ids)
    db.delete(question)
    db.commit()
    return RedirectResponse(url="/teacher/questions", status_code=303)


@router.get("/stats")
def teacher_stats(request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    question_ids = [
        q.id for q in db.query(Question).filter(Question.created_by == user_id).all()
    ]
    total_questions = len(question_ids)
    total_records = db.query(Record).filter(Record.question_id.in_(question_ids)).count() if question_ids else 0
    correct_records = (
        db.query(Record)
        .filter(Record.question_id.in_(question_ids), Record.is_correct == True)
        .count()
    ) if question_ids else 0
    accuracy = round(correct_records / total_records * 100, 1) if total_records > 0 else 0

    per_question = (
        db.query(
            Question.id,
            Question.content,
            Question.subject,
            Question.q_type,
            sa_func.count(Record.id).label("attempts"),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
        )
        .outerjoin(Record, Record.question_id == Question.id)
        .filter(Question.created_by == user_id)
        .group_by(Question.id)
        .order_by(sa_func.count(Record.id).desc())
        .limit(20)
        .all()
    )

    # Get student stats - filter by class membership for regular teachers
    if question_ids:
        if is_admin(db, user_id):
            # Admin can see all students
            student_records = (
                db.query(
                    User.id,
                    User.username,
                    User.display_name,
                    sa_func.count(Record.id).label("total"),
                    sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
                )
                .join(Record, Record.user_id == User.id)
                .filter(Record.question_id.in_(question_ids), User.role == "student")
                .group_by(User.id)
                .all()
            )
        else:
            # Regular teacher: only students from their classes
            teacher_class_ids = [c.id for c in db.query(ClassGroup).filter(ClassGroup.created_by == user_id).all()]
            if teacher_class_ids:
                student_ids_in_classes = [m.user_id for m in db.query(ClassMember).filter(ClassMember.class_id.in_(teacher_class_ids)).all()]
                student_records = (
                    db.query(
                        User.id,
                        User.username,
                        User.display_name,
                        sa_func.count(Record.id).label("total"),
                        sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
                    )
                    .join(Record, Record.user_id == User.id)
                    .filter(
                        Record.question_id.in_(question_ids),
                        User.role == "student",
                        User.id.in_(student_ids_in_classes)
                    )
                    .group_by(User.id)
                    .all()
                )
            else:
                student_records = []
    else:
        student_records = []
    student_stats = [
        {
            "id": s.id,
            "username": s.username,
            "display_name": s.display_name,
            "total": s.total,
            "correct": int(s.correct or 0),
            "accuracy": round((s.correct or 0) / s.total * 100, 1),
        }
        for s in student_records
    ]
    student_stats.sort(key=lambda x: x["accuracy"])

    return request.app.state.templates.TemplateResponse(
        "teacher/stats.html",
        {
            "request": request,
            "total_questions": total_questions,
            "total_records": total_records,
            "accuracy": accuracy,
            "per_question": per_question,
            "student_stats": student_stats,
            "question_types": QUESTION_TYPES,
        },
    )


@router.get("/students/{student_id}")
def student_detail(student_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    student = db.query(User).filter(User.id == student_id, User.role == "student").first()
    if not student:
        raise HTTPException(status_code=404, detail="学生不存在")
    if not teacher_owns_student(db, user_id, student_id):
        raise HTTPException(status_code=404)

    total = db.query(Record).filter(Record.user_id == student_id).count()
    correct = db.query(Record).filter(Record.user_id == student_id, Record.is_correct == True).count()
    accuracy = round(correct / total * 100, 1) if total > 0 else 0

    by_subject = (
        db.query(
            Question.subject,
            sa_func.count(Record.id),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)),
        )
        .join(Question, Record.question_id == Question.id)
        .filter(Record.user_id == student_id)
        .group_by(Question.subject)
        .all()
    )
    subject_stats = [
        {"name": s, "total": t, "correct": c or 0, "accuracy": round((c or 0) / t * 100, 1) if t > 0 else 0}
        for s, t, c in by_subject
    ]
    subject_stats.sort(key=lambda x: x["accuracy"])

    by_type = (
        db.query(
            Question.q_type,
            sa_func.count(Record.id),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)),
        )
        .join(Question, Record.question_id == Question.id)
        .filter(Record.user_id == student_id)
        .group_by(Question.q_type)
        .all()
    )
    type_stats = [
        {"type": QUESTION_TYPES.get(qt, qt), "total": t, "correct": c or 0, "accuracy": round((c or 0) / t * 100, 1) if t > 0 else 0}
        for qt, t, c in by_type
    ]

    wrong_qids = set(
        r[0] for r in db.query(Record.question_id)
        .filter(Record.user_id == student_id, Record.is_correct == False)
        .distinct().all()
    )
    mistake_count = len(wrong_qids)
    if wrong_qids:
        corrected = set(
            r[0] for r in db.query(Record.question_id)
            .filter(Record.user_id == student_id, Record.question_id.in_(wrong_qids), Record.is_correct == True)
            .distinct().all()
        )
        mistake_count = len(wrong_qids - corrected)

    daily = (
        db.query(
            sa_func.date(Record.created_at).label("date"),
            sa_func.count(Record.id),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)),
        )
        .filter(Record.user_id == student_id)
        .group_by(sa_func.date(Record.created_at))
        .order_by(sa_func.date(Record.created_at).desc())
        .limit(14)
        .all()
    )
    daily_stats = [
        {"date": str(d), "total": t, "correct": c or 0, "accuracy": round((c or 0) / t * 100, 1) if t > 0 else 0}
        for d, t, c in daily
    ]

    return request.app.state.templates.TemplateResponse(
        "teacher/student_detail.html",
        {
            "request": request,
            "student": student,
            "total": total,
            "correct": correct,
            "accuracy": accuracy,
            "subject_stats": subject_stats,
            "type_stats": type_stats,
            "mistake_count": mistake_count,
            "daily_stats": daily_stats,
            "question_types": QUESTION_TYPES,
        },
    )


@router.get("/stats/export/pdf")
def export_stats_pdf(request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    user = db.query(User).filter(User.id == user_id).first()
    question_ids = [q.id for q in db.query(Question).filter(Question.created_by == user_id).all()]
    total_questions = len(question_ids)
    total_records = db.query(Record).filter(Record.question_id.in_(question_ids)).count() if question_ids else 0
    correct_records = db.query(Record).filter(Record.question_id.in_(question_ids), Record.is_correct == True).count() if question_ids else 0
    accuracy = round(correct_records / total_records * 100, 1) if total_records > 0 else 0

    if is_admin(db, user_id):
        students = db.query(User).filter(User.role == "student").all()
    else:
        teacher_class_ids = [c.id for c in db.query(ClassGroup).filter(ClassGroup.created_by == user_id).all()]
        if teacher_class_ids:
            student_ids = [m.user_id for m in db.query(ClassMember).filter(ClassMember.class_id.in_(teacher_class_ids)).all()]
            students = db.query(User).filter(User.id.in_(student_ids), User.role == "student").all() if student_ids else []
        else:
            students = []

    # 批量查询所有学生的统计数据，避免 N+1 查询
    student_stats = []
    if students and question_ids:
        student_id_list = [s.id for s in students]
        # 一次性查询所有学生的记录统计
        stats_query = (
            db.query(
                Record.user_id,
                sa_func.count(Record.id).label("total"),
                sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct")
            )
            .filter(Record.user_id.in_(student_id_list), Record.question_id.in_(question_ids))
            .group_by(Record.user_id)
            .all()
        )
        stats_map = {row.user_id: (row.total, row.correct or 0) for row in stats_query}

        for s in students:
            if s.id in stats_map:
                s_total, s_correct = stats_map[s.id]
                student_stats.append({
                    "username": s.username,
                    "display_name": s.display_name,
                    "total": s_total,
                    "correct": s_correct,
                    "accuracy": round(s_correct / s_total * 100, 1)
                })
    student_stats.sort(key=lambda x: x["accuracy"])

    from app.utils.report import generate_teacher_report
    buffer = generate_teacher_report(user.display_name or user.username, total_questions, total_records, accuracy, [], student_stats)

    return Response(
        content=buffer.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=teacher_report.pdf"},
    )


@router.get("/students/{student_id}/parent-report")
def parent_report(student_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    student = db.query(User).filter(User.id == student_id, User.role == "student").first()
    if not student:
        raise HTTPException(status_code=404, detail="学生不存在")
    if not teacher_owns_student(db, user_id, student_id):
        raise HTTPException(status_code=404)

    today = datetime.now()
    monday = today - timedelta(days=today.weekday())
    week_start = datetime(monday.year, monday.month, monday.day)

    week_records = db.query(Record).filter(
        Record.user_id == student_id,
        Record.created_at >= week_start,
    ).all()

    week_total = len(week_records)
    week_correct = sum(1 for r in week_records if r.is_correct)
    week_accuracy = round(week_correct / week_total * 100, 1) if week_total > 0 else 0

    week_question_ids = [r.question_id for r in week_records]
    chapter_stats = {}
    if week_question_ids:
        chapter_rows = (
            db.query(
                Question.subject,
                Question.chapter,
                sa_func.count(Record.id).label("total"),
                sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
            )
            .join(Question, Record.question_id == Question.id)
            .filter(Record.user_id == student_id, Record.created_at >= week_start)
            .group_by(Question.subject, Question.chapter)
            .all()
        )
        for row in chapter_rows:
            acc = round((row.correct or 0) / row.total * 100, 1) if row.total > 0 else 0
            chapter_stats[f"{row.subject}-{row.chapter or '未分类'}"] = {
                "subject": row.subject,
                "chapter": row.chapter or "未分类",
                "total": row.total,
                "correct": int(row.correct or 0),
                "accuracy": acc,
            }

    weak_points = [v for v in chapter_stats.values() if v["accuracy"] < 60]
    weak_points.sort(key=lambda x: x["accuracy"])

    suggestions = []
    for wp in weak_points:
        available = db.query(Question).filter(
            Question.subject == wp["subject"],
            Question.chapter == wp["chapter"] if wp["chapter"] != "未分类" else "",
        ).count()
        suggestions.append({
            "subject": wp["subject"],
            "chapter": wp["chapter"],
            "accuracy": wp["accuracy"],
            "available_questions": available,
        })

    return request.app.state.templates.TemplateResponse(
        "teacher/parent_report.html",
        {
            "request": request,
            "student": student,
            "week_total": week_total,
            "week_correct": week_correct,
            "week_accuracy": week_accuracy,
            "weak_points": weak_points,
            "suggestions": suggestions,
            "week_start": week_start.strftime("%Y-%m-%d"),
            "today": today.strftime("%Y-%m-%d"),
        },
    )


@router.get("/students/{student_id}/parent-report/pdf")
def parent_report_pdf(student_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    student = db.query(User).filter(User.id == student_id, User.role == "student").first()
    if not student:
        raise HTTPException(status_code=404, detail="学生不存在")
    if not teacher_owns_student(db, user_id, student_id):
        raise HTTPException(status_code=404)

    today = datetime.now()
    monday = today - timedelta(days=today.weekday())
    week_start = datetime(monday.year, monday.month, monday.day)

    week_records = db.query(Record).filter(
        Record.user_id == student_id,
        Record.created_at >= week_start,
    ).all()

    week_total = len(week_records)
    week_correct = sum(1 for r in week_records if r.is_correct)
    week_accuracy = round(week_correct / week_total * 100, 1) if week_total > 0 else 0

    chapter_stats = {}
    chapter_rows = (
        db.query(
            Question.subject,
            Question.chapter,
            sa_func.count(Record.id).label("total"),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
        )
        .join(Question, Record.question_id == Question.id)
        .filter(Record.user_id == student_id, Record.created_at >= week_start)
        .group_by(Question.subject, Question.chapter)
        .all()
    )
    for row in chapter_rows:
        acc = round((row.correct or 0) / row.total * 100, 1) if row.total > 0 else 0
        chapter_stats[f"{row.subject}-{row.chapter or '未分类'}"] = {
            "subject": row.subject,
            "chapter": row.chapter or "未分类",
            "total": row.total,
            "correct": int(row.correct or 0),
            "accuracy": acc,
        }

    weak_points = [v for v in chapter_stats.values() if v["accuracy"] < 60]
    weak_points.sort(key=lambda x: x["accuracy"])

    suggestions = []
    for wp in weak_points:
        available = db.query(Question).filter(
            Question.subject == wp["subject"],
            Question.chapter == wp["chapter"] if wp["chapter"] != "未分类" else "",
        ).count()
        suggestions.append({
            "subject": wp["subject"],
            "chapter": wp["chapter"],
            "accuracy": wp["accuracy"],
            "available_questions": available,
        })

    from app.utils.report import generate_parent_report
    buffer = generate_parent_report(
        student.display_name or student.username,
        week_total,
        week_correct,
        week_accuracy,
        weak_points,
        suggestions,
        week_start.strftime("%Y-%m-%d"),
        today.strftime("%Y-%m-%d"),
    )

    return Response(
        content=buffer.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=parent_report_{student_id}.pdf"},
    )


@router.get("/students/{student_id}/export/pdf")
def export_student_pdf(student_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    student = db.query(User).filter(User.id == student_id, User.role == "student").first()
    if not student:
        raise HTTPException(status_code=404, detail="学生不存在")
    if not teacher_owns_student(db, user_id, student_id):
        raise HTTPException(status_code=404)

    total = db.query(Record).filter(Record.user_id == student_id).count()
    correct = db.query(Record).filter(Record.user_id == student_id, Record.is_correct == True).count()
    accuracy = round(correct / total * 100, 1) if total > 0 else 0

    by_subject = (
        db.query(Question.subject, sa_func.count(Record.id), sa_func.sum(sa_func.cast(Record.is_correct, Integer)))
        .join(Question, Record.question_id == Question.id)
        .filter(Record.user_id == student_id)
        .group_by(Question.subject).all()
    )
    subject_stats = [{"name": s, "total": t, "correct": c or 0, "accuracy": round((c or 0) / t * 100, 1) if t > 0 else 0} for s, t, c in by_subject]

    from app.utils.report import generate_student_report
    buffer = generate_student_report(student.display_name or student.username, total, correct, accuracy, subject_stats, [])

    return Response(
        content=buffer.read(),
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=student_{student_id}_report.pdf"},
    )


@router.get("/students")
def student_management(request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    classes = db.query(ClassGroup).filter(ClassGroup.created_by == user_id).order_by(ClassGroup.name).all()

    class_ids = [c.id for c in classes]
    member_rows = (
        db.query(
            ClassMember.class_id,
            User.id,
            User.username,
            User.display_name,
            User.is_guest,
            sa_func.count(Record.id).label("total"),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
        )
        .join(User, User.id == ClassMember.user_id)
        .outerjoin(Record, Record.user_id == User.id)
        .filter(ClassMember.class_id.in_(class_ids))
        .group_by(ClassMember.class_id, User.id)
        .all()
    )
    class_member_map = {}
    class_member_count = {}
    for row in member_rows:
        class_member_map.setdefault(row.class_id, []).append({
            "id": row.id,
            "username": row.username,
            "display_name": row.display_name,
            "is_guest": row.is_guest,
            "total": row.total,
            "correct": int(row.correct or 0),
            "accuracy": round((row.correct or 0) / row.total * 100, 1) if row.total > 0 else 0,
        })
        class_member_count[row.class_id] = class_member_count.get(row.class_id, 0) + 1

    MEMBERS_PER_CLASS = 20
    class_data = []
    for cls in classes:
        all_members = class_member_map.get(cls.id, [])
        shown_members = all_members[:MEMBERS_PER_CLASS]
        class_data.append({
            "id": cls.id,
            "name": cls.name,
            "members": shown_members,
            "total_members": len(all_members),
            "has_more": len(all_members) > MEMBERS_PER_CLASS,
        })

    # Get guest students - filter by pending join requests to teacher's classes
    if is_admin(db, user_id):
        # Admin can see all guest students
        guests = db.query(User).filter(User.role == "student", User.is_guest == True).all()
    else:
        # Regular teacher: only guests who have pending requests to their classes
        if class_ids:
            guest_user_ids = [
                jr.user_id for jr in db.query(ClassJoinRequest)
                .filter(ClassJoinRequest.class_id.in_(class_ids), ClassJoinRequest.status == "pending")
                .all()
            ]
            guests = db.query(User).filter(
                User.id.in_(guest_user_ids),
                User.role == "student",
                User.is_guest == True
            ).all() if guest_user_ids else []
        else:
            guests = []
    guest_data = [
        {"id": g.id, "username": g.username, "display_name": g.display_name, "guest_expires_at": g.guest_expires_at}
        for g in guests
    ]

    pending_join_requests = (
        db.query(ClassJoinRequest)
        .filter(ClassJoinRequest.class_id.in_(class_ids), ClassJoinRequest.status == "pending")
        .all()
    ) if class_ids else []
    join_request_data = []
    for jr in pending_join_requests:
        jr_user = db.query(User).filter(User.id == jr.user_id).first()
        jr_class = db.query(ClassGroup).filter(ClassGroup.id == jr.class_id).first()
        if jr_user and jr_class:
            join_request_data.append({
                "id": jr.id,
                "user_id": jr_user.id,
                "username": jr_user.username,
                "display_name": jr.display_name or jr_user.display_name,
                "class_name": jr_class.name,
                "class_id": jr_class.id,
                "created_at": jr.created_at,
            })

    from app.models import AccountRecoveryRequest
    class_recovery_requests = (
        db.query(AccountRecoveryRequest)
        .filter(AccountRecoveryRequest.class_id.in_(class_ids), AccountRecoveryRequest.status == "pending")
        .all()
    ) if class_ids else []
    recovery_data = []
    for rr in class_recovery_requests:
        rr_class = db.query(ClassGroup).filter(ClassGroup.id == rr.class_id).first()
        recovery_data.append({
            "id": rr.id,
            "username": rr.username,
            "display_name": rr.display_name,
            "class_name": rr_class.name if rr_class else "",
            "created_at": rr.created_at,
        })

    return request.app.state.templates.TemplateResponse(
        "teacher/students.html",
        {"request": request, "class_data": class_data, "guest_data": guest_data, "join_request_data": join_request_data, "recovery_data": recovery_data, "csrf_token": request.session.get("csrf_token", "")},
    )


@router.post("/students/{student_id}/approve")
async def approve_student(student_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    form = await request.form()
    class_id = form.get("class_id", "").strip()
    student = db.query(User).filter(User.id == student_id, User.role == "student").first()
    if not student:
        return RedirectResponse(url="/teacher/students", status_code=303)
    if not student.is_guest:
        return RedirectResponse(url="/teacher/students", status_code=303)
    if not class_id:
        return RedirectResponse(url="/teacher/students", status_code=303)
    class_id_value = parse_int(class_id, min_value=1)
    if class_id_value is None:
        return RedirectResponse(url="/teacher/students", status_code=303)
    cls = db.query(ClassGroup).filter(ClassGroup.id == class_id_value, ClassGroup.created_by == user_id).first()
    if cls:
        if student.class_id and student.class_id != cls.id:
            return RedirectResponse(url="/teacher/students?error=student_in_other_class", status_code=303)

        existing = db.query(ClassMember).filter(ClassMember.class_id == cls.id, ClassMember.user_id == student.id).first()
        if not existing:
            db.add(ClassMember(class_id=cls.id, user_id=student.id))
        student.is_guest = False
        student.guest_expires_at = None
        student.class_id = cls.id
        db.commit()
        db.add(Notification(
            user_id=student.id,
            title="审核通过",
            content=f"您已被教师审核通过，正式加入班级，现在可以正常使用所有功能。",
        ))
        db.commit()
    return RedirectResponse(url="/teacher/students", status_code=303)


@router.post("/join-requests/{request_id}/approve")
async def approve_join_request(request_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    join_req = db.query(ClassJoinRequest).filter(ClassJoinRequest.id == request_id, ClassJoinRequest.status == "pending").first()
    if not join_req:
        return RedirectResponse(url="/teacher/students", status_code=303)
    cls = db.query(ClassGroup).filter(ClassGroup.id == join_req.class_id, ClassGroup.created_by == user_id).first()
    if not cls:
        return RedirectResponse(url="/teacher/students", status_code=303)
    student = db.query(User).filter(User.id == join_req.user_id).first()
    if not student:
        return RedirectResponse(url="/teacher/students", status_code=303)

    if student.class_id and student.class_id != cls.id:
        return RedirectResponse(url="/teacher/students?error=student_in_other_class", status_code=303)

    existing_member = db.query(ClassMember).filter(ClassMember.class_id == cls.id, ClassMember.user_id == student.id).first()
    if not existing_member:
        db.add(ClassMember(class_id=cls.id, user_id=student.id))
    student.is_guest = False
    student.guest_expires_at = None
    student.class_id = cls.id
    student.join_mode = "formal"
    join_req.status = "approved"
    join_req.reviewed_by = user_id
    from datetime import datetime
    join_req.reviewed_at = datetime.now()
    db.add(Notification(
        user_id=student.id,
        title="入班申请已通过",
        content=f"您申请加入班级「{cls.name}」已通过审核，现在可以正常使用所有功能。",
    ))
    db.commit()
    return RedirectResponse(url="/teacher/students", status_code=303)


@router.post("/join-requests/{request_id}/reject")
async def reject_join_request(request_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    join_req = db.query(ClassJoinRequest).filter(ClassJoinRequest.id == request_id, ClassJoinRequest.status == "pending").first()
    if not join_req:
        return RedirectResponse(url="/teacher/students", status_code=303)
    cls = db.query(ClassGroup).filter(ClassGroup.id == join_req.class_id, ClassGroup.created_by == user_id).first()
    if not cls:
        return RedirectResponse(url="/teacher/students", status_code=303)
    join_req.status = "rejected"
    join_req.reviewed_by = user_id
    from datetime import datetime
    join_req.reviewed_at = datetime.now()
    student = db.query(User).filter(User.id == join_req.user_id).first()
    if student:
        db.add(Notification(
            user_id=student.id,
            title="入班申请未通过",
            content=f"您申请加入班级「{cls.name}」未通过审核，请选择其他班级或以游客身份体验。",
        ))
    db.commit()
    return RedirectResponse(url="/teacher/students", status_code=303)


@router.get("/invite")
def invite_manager(request: Request, db: Annotated[Session, Depends(get_db)]):
    require_admin(request, db)
    config = db.query(SiteConfig).filter(SiteConfig.key == "teacher_invite_code").first()
    current_codes = config.value if config else ""
    return request.app.state.templates.TemplateResponse(
        "teacher/invite.html",
        {"request": request, "current_codes": current_codes, "csrf_token": request.session.get("csrf_token", "")},
    )


@router.post("/invite/update")
async def update_invite(request: Request, db: Annotated[Session, Depends(get_db)]):
    require_admin(request, db)
    await validate_csrf_async(request)
    form = await request.form()
    codes = sanitize_input(form.get("codes", "").strip(), max_length=500)
    config = db.query(SiteConfig).filter(SiteConfig.key == "teacher_invite_code").first()
    if not config:
        config = SiteConfig(key="teacher_invite_code", value=codes)
        db.add(config)
    else:
        config.value = codes
    db.commit()
    return RedirectResponse(url="/teacher/invite", status_code=303)


@router.get("/classes/{class_id}/import-students")
def import_students_page(class_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    cls = db.query(ClassGroup).filter(ClassGroup.id == class_id, ClassGroup.created_by == user_id).first()
    if not cls:
        raise HTTPException(status_code=404, detail="班级不存在")
    return request.app.state.templates.TemplateResponse(
        "teacher/student_import.html",
        {
            "request": request,
            "class_info": cls,
            "error": None,
            "success": None,
            "created_count": 0,
            "skipped_count": 0,
            "csrf_token": request.session.get("csrf_token", ""),
        },
    )


@router.post("/classes/{class_id}/import-students")
async def import_students(class_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    await validate_csrf_async(request)
    cls = db.query(ClassGroup).filter(ClassGroup.id == class_id, ClassGroup.created_by == user_id).first()
    if not cls:
        raise HTTPException(status_code=404, detail="班级不存在")

    form = await request.form()
    csv_text = form.get("csv_text", "").strip()

    if not csv_text:
        return request.app.state.templates.TemplateResponse(
            "teacher/student_import.html",
            {
                "request": request,
                "class_info": cls,
                "error": "请输入CSV数据",
                "success": None,
                "created_count": 0,
                "skipped_count": 0,
                "csrf_token": request.session.get("csrf_token", ""),
            },
        )

    created_count = 0
    updated_count = 0
    skipped_count = 0
    skipped_details = []

    try:
        reader = csv.DictReader(io.StringIO(csv_text))
        for row in reader:
            username = row.get("username", "").strip()
            display_name = row.get("display_name", "").strip()
            if not username:
                continue

            existing = db.query(User).filter(User.username == username).first()
            if existing:
                if existing.class_id and existing.class_id != cls.id:
                    skipped_count += 1
                    skipped_details.append({
                        "username": username,
                        "reason": f"已在其他班级（class_id={existing.class_id}）"
                    })
                    continue

                member = db.query(ClassMember).filter(
                    ClassMember.class_id == cls.id,
                    ClassMember.user_id == existing.id,
                ).first()
                if not member:
                    db.add(ClassMember(class_id=cls.id, user_id=existing.id))
                existing.class_id = cls.id
                if existing.is_guest:
                    existing.is_guest = False
                    existing.guest_expires_at = None
                updated_count += 1
            else:
                new_user = User(
                    username=username,
                    password_hash=User.hash_password("abc12345"),
                    role="student",
                    display_name=display_name or username,
                    class_id=cls.id,
                )
                db.add(new_user)
                db.flush()
                db.add(ClassMember(class_id=cls.id, user_id=new_user.id))
                created_count += 1

        db.commit()
    except Exception as e:
        return request.app.state.templates.TemplateResponse(
            "teacher/student_import.html",
            {
                "request": request,
                "class_info": cls,
                "error": "导入失败，请检查文件格式是否正确",
                "success": None,
                "created_count": 0,
                "updated_count": 0,
                "skipped_count": 0,
                "csrf_token": request.session.get("csrf_token", ""),
            },
        )

    return request.app.state.templates.TemplateResponse(
        "teacher/student_import.html",
        {
            "request": request,
            "class_info": cls,
            "error": None,
            "success": f"导入完成：新建 {created_count} 人，更新 {updated_count} 人，跳过 {skipped_count} 人",
            "created_count": created_count,
            "updated_count": updated_count,
            "skipped_count": skipped_count,
            "skipped_details": skipped_details,
            "csrf_token": request.session.get("csrf_token", ""),
        },
    )


@router.get("/classes/{class_id}/stats")
def class_stats(class_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    cls = db.query(ClassGroup).filter(ClassGroup.id == class_id, ClassGroup.created_by == user_id).first()
    if not cls:
        raise HTTPException(status_code=404, detail="班级不存在")

    member_ids = [m.user_id for m in db.query(ClassMember).filter(ClassMember.class_id == class_id).all()]

    seven_days_ago = datetime.now() - timedelta(days=7)
    accuracy_trend = (
        db.query(
            sa_func.date(Record.created_at).label("date"),
            sa_func.count(Record.id).label("total"),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
        )
        .filter(Record.user_id.in_(member_ids), Record.created_at >= seven_days_ago)
        .group_by(sa_func.date(Record.created_at))
        .order_by(sa_func.date(Record.created_at))
        .all()
    )
    trend_data = [
        {"date": str(row.date), "total": row.total, "correct": int(row.correct or 0), "accuracy": round((row.correct or 0) / row.total * 100, 1) if row.total > 0 else 0}
        for row in accuracy_trend
    ]

    weak_chapters = (
        db.query(
            Question.subject,
            Question.chapter,
            sa_func.count(Record.id).label("total"),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
        )
        .join(Question, Record.question_id == Question.id)
        .filter(Record.user_id.in_(member_ids))
        .group_by(Question.subject, Question.chapter)
        .all()
    )
    chapter_data = [
        {"subject": row.subject, "chapter": row.chapter, "total": row.total, "correct": int(row.correct or 0), "accuracy": round((row.correct or 0) / row.total * 100, 1) if row.total > 0 else 0}
        for row in weak_chapters
    ]
    chapter_data.sort(key=lambda x: x["accuracy"])
    weak_chapters_top5 = chapter_data[:5]

    student_ranking = (
        db.query(
            User.id,
            User.username,
            User.display_name,
            sa_func.count(Record.id).label("total"),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
        )
        .join(Record, Record.user_id == User.id)
        .filter(Record.user_id.in_(member_ids))
        .group_by(User.id)
        .all()
    )
    ranking = [
        {"id": s.id, "username": s.username, "display_name": s.display_name, "total": s.total, "correct": int(s.correct or 0), "accuracy": round((s.correct or 0) / s.total * 100, 1) if s.total > 0 else 0}
        for s in student_ranking
    ]
    ranking.sort(key=lambda x: x["accuracy"], reverse=True)

    progress_list = []
    for s in student_ranking:
        records = db.query(Record).filter(Record.user_id == s.id).order_by(Record.created_at).all()
        if len(records) < 2:
            continue
        mid = len(records) // 2
        first_half = records[:mid]
        second_half = records[mid:]
        first_correct = sum(1 for r in first_half if r.is_correct)
        first_acc = first_correct / len(first_half) * 100 if first_half else 0
        second_correct = sum(1 for r in second_half if r.is_correct)
        second_acc = second_correct / len(second_half) * 100 if second_half else 0
        improvement = round(second_acc - first_acc, 1)
        progress_list.append({
            "id": s.id,
            "username": s.username,
            "display_name": s.display_name,
            "first_acc": round(first_acc, 1),
            "second_acc": round(second_acc, 1),
            "improvement": improvement,
        })
    progress_list.sort(key=lambda x: x["improvement"], reverse=True)
    progress_top5 = progress_list[:5]

    all_member_users = (
        db.query(User.id, User.username, User.display_name)
        .join(ClassMember, ClassMember.user_id == User.id)
        .filter(ClassMember.class_id == class_id)
        .all()
    )
    practiced_ids = set(s.id for s in student_ranking)
    no_practice = [
        {"id": u.id, "username": u.username, "display_name": u.display_name}
        for u in all_member_users if u.id not in practiced_ids
    ]

    return request.app.state.templates.TemplateResponse(
        "teacher/class_stats.html",
        {
            "request": request,
            "class_info": cls,
            "trend_data": trend_data,
            "weak_chapters": weak_chapters_top5,
            "ranking": ranking,
            "progress_top5": progress_top5,
            "no_practice": no_practice,
        },
    )


@router.get("/classes/{class_id}/export/excel")
def export_class_excel(class_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    cls = db.query(ClassGroup).filter(ClassGroup.id == class_id, ClassGroup.created_by == user_id).first()
    if not cls:
        raise HTTPException(status_code=404, detail="班级不存在")

    member_ids = [m.user_id for m in db.query(ClassMember).filter(ClassMember.class_id == class_id).all()]

    student_ranking = (
        db.query(
            User.id,
            User.username,
            User.display_name,
            sa_func.count(Record.id).label("total"),
            sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct"),
        )
        .join(Record, Record.user_id == User.id)
        .filter(Record.user_id.in_(member_ids))
        .group_by(User.id)
        .all()
    ) if member_ids else []

    ranking = [
        {"id": s.id, "username": s.username, "display_name": s.display_name, "total": s.total, "correct": int(s.correct or 0), "accuracy": round((s.correct or 0) / s.total * 100, 1) if s.total > 0 else 0}
        for s in student_ranking
    ]
    ranking.sort(key=lambda x: x["accuracy"], reverse=True)

    practiced_ids = set(s.id for s in student_ranking)
    no_practice_users = db.query(User).filter(User.id.in_(member_ids), User.role == "student").all() if member_ids else []
    no_practice = [u for u in no_practice_users if u.id not in practiced_ids]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "班级报告"
    ws.append(["排名", "姓名", "用户名", "做题数", "正确数", "正确率"])
    for idx, s in enumerate(ranking, 1):
        ws.append([idx, s["display_name"] or s["username"], s["username"], s["total"], s["correct"], f'{s["accuracy"]}%'])
    for u in no_practice:
        ws.append(["-", u.display_name or u.username, u.username, 0, 0, "0%"])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = urllib.parse.quote(f"{cls.name}_班级报告.xlsx")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})


@router.get("/assignments/{assignment_id}/export/excel")
def export_assignment_excel(assignment_id: int, request: Request, db: Annotated[Session, Depends(get_db)]):
    user_id = require_teacher(request, db)
    assignment = db.query(Assignment).filter(Assignment.id == assignment_id, Assignment.created_by == user_id).first()
    if not assignment:
        raise HTTPException(status_code=404, detail="作业不存在")

    qid_list = [int(x.strip()) for x in assignment.question_ids.split(",") if x.strip().isdigit()]

    completed_records = db.query(AssignmentRecord).filter(
        AssignmentRecord.assignment_id == assignment_id,
        AssignmentRecord.completed == True,
    ).all()
    completed_user_ids = {r.user_id for r in completed_records}

    if assignment.class_id:
        members = db.query(ClassMember).filter(ClassMember.class_id == assignment.class_id).all()
        student_ids = [m.user_id for m in members]
    else:
        classes = db.query(ClassGroup).filter(ClassGroup.created_by == user_id).all()
        class_ids = [c.id for c in classes]
        members = db.query(ClassMember).filter(ClassMember.class_id.in_(class_ids)).all() if class_ids else []
        student_ids = list({m.user_id for m in members})

    students = db.query(User).filter(User.id.in_(student_ids), User.role == "student").all() if student_ids else []

    # 批量查询所有学生的作业统计，避免 N+1 查询
    stats_map = {}
    if students and qid_list:
        student_id_list = [s.id for s in students]
        stats_query = (
            db.query(
                Record.user_id,
                sa_func.count(Record.id).label("total"),
                sa_func.sum(sa_func.cast(Record.is_correct, Integer)).label("correct")
            )
            .filter(Record.user_id.in_(student_id_list), Record.question_id.in_(qid_list))
            .group_by(Record.user_id)
            .all()
        )
        stats_map = {row.user_id: (row.total, row.correct or 0) for row in stats_query}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "作业完成情况"
    ws.append(["姓名", "用户名", "完成状态", "正确率"])

    for s in students:
        status = "已完成" if s.id in completed_user_ids else "未完成"
        accuracy = 0.0
        if s.id in completed_user_ids and s.id in stats_map:
            total, correct = stats_map[s.id]
            accuracy = round(correct / total * 100, 1) if total > 0 else 0.0
        ws.append([s.display_name or s.username, s.username, status, f"{accuracy}%"])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = urllib.parse.quote(f"{assignment.title}_作业完成情况.xlsx")
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"})
